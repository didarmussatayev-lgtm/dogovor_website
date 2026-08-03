from __future__ import annotations

import base64
import logging
import re
import subprocess
from datetime import date
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docxtpl import DocxTemplate, InlineImage
from docx.shared import Mm
from jinja2 import TemplateSyntaxError

logger = logging.getLogger(__name__)

_JINJA_PRINT_RE = re.compile(r"\{\{.*?\}\}", flags=re.DOTALL)
_XML_TAG_RE = re.compile(r"<[^>]+>")
_SAFE_JINJA_KEY_RE = re.compile(r"^[A-Za-z_][A-Za-z0-9_]*$")
_PLACEHOLDER_ALIASES = {
    "дата рождения": "birth_date",
    "пол": "gender",
}


def _decode_signature(signature_base64: str) -> bytes:
    """Strip optional data-URL prefix and return raw PNG bytes."""
    if "," in signature_base64:
        signature_base64 = signature_base64.split(",", 1)[1]
    return base64.b64decode(signature_base64)


def _normalize_template_expression(raw_expression: str) -> tuple[str, str | None]:
    plain_expression = _XML_TAG_RE.sub("", raw_expression)
    if not plain_expression.startswith("{{") or not plain_expression.endswith("}}"):
        return raw_expression, None

    placeholder = re.sub(r"\s+", " ", plain_expression[2:-2]).strip()
    alias_key = placeholder.lower()
    mapped = _PLACEHOLDER_ALIASES.get(alias_key)

    if mapped and mapped != placeholder:
        return f"{{{{ {mapped} }}}}", placeholder

    if _SAFE_JINJA_KEY_RE.fullmatch(placeholder):
        return raw_expression, None

    return raw_expression, placeholder


def _prepare_template_for_render(template_path: str | Path, output_dir: str | Path, output_basename: str) -> tuple[Path, list[str]]:
    source = Path(template_path)
    normalized_path = Path(output_dir) / f"{output_basename}_template.docx"
    rewritten_placeholders: list[str] = []
    suspicious_placeholders: list[str] = []

    with ZipFile(source, "r") as src, ZipFile(normalized_path, "w", compression=ZIP_DEFLATED) as dst:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename.startswith("word/") and info.filename.endswith(".xml"):
                xml = data.decode("utf-8")

                def _replace(match: re.Match[str]) -> str:
                    rewritten, suspicious = _normalize_template_expression(match.group(0))
                    if rewritten != match.group(0):
                        rewritten_placeholders.append(suspicious or "")
                    elif suspicious:
                        suspicious_placeholders.append(suspicious)
                    return rewritten

                xml = _JINJA_PRINT_RE.sub(_replace, xml)
                data = xml.encode("utf-8")
            dst.writestr(info, data)

    if rewritten_placeholders:
        cleaned = sorted({p for p in rewritten_placeholders if p})
        logger.info("Normalized legacy placeholders in %s: %s", source.name, ", ".join(cleaned))

    return normalized_path, sorted(set(suspicious_placeholders))


def generate_docx(
    template_path: str | Path,
    full_name: str,
    phone: str,
    iin: str,
    birth_date: str,
    gender_display: str,
    allergy: str,
    procedure: str,
    signature_base64: str,
    agreement_id: str,
    output_basename: str,
    output_dir: str | Path,
) -> Path:
    """Fill the DOCX template and return the path to the generated file."""
    normalized_template_path, suspicious_placeholders = _prepare_template_for_render(
        template_path=template_path,
        output_dir=output_dir,
        output_basename=output_basename,
    )
    tpl = DocxTemplate(normalized_template_path)

    # Decode signature and write to a temp PNG so InlineImage can read it
    sig_bytes = _decode_signature(signature_base64)
    sig_tmp = Path(output_dir) / f"{agreement_id}_sig.png"
    sig_tmp.write_bytes(sig_bytes)

    context = {
        "full_name": full_name,
        "phone": phone,
        "iin": iin,
        "birth_date": birth_date,
        "gender": gender_display,
        "allergy": allergy,
        "procedure": procedure,
        "date": date.today().strftime("%d.%m.%Y"),
        "agreement_id": agreement_id,
        "signature": InlineImage(tpl, str(sig_tmp), width=Mm(50)),
    }

    try:
        tpl.render(context)
    except TemplateSyntaxError as exc:
        hint = ""
        if suspicious_placeholders:
            hint = (
                f" Likely invalid placeholder(s): {', '.join(suspicious_placeholders[:3])}. "
                "Use ASCII underscore keys like {{ birth_date }} or {{ gender }}."
            )
        raise RuntimeError(f"Template syntax error in {Path(template_path).name}: {exc}.{hint}") from exc

    docx_path = Path(output_dir) / f"{output_basename}.docx"
    tpl.save(str(docx_path))
    logger.info("DOCX generated: %s", docx_path)
    return docx_path


def convert_to_pdf(docx_path: Path, output_dir: Path) -> Path:
    """Convert a DOCX file to PDF using LibreOffice headless."""
    cmd = [
        "libreoffice",
        "--headless",
        "--convert-to",
        "pdf",
        "--outdir",
        str(output_dir),
        str(docx_path),
    ]
    logger.info("Running LibreOffice: %s", " ".join(cmd))
    try:
        result = subprocess.run(
            cmd,
            capture_output=True,
            text=True,
            timeout=120,
        )
    except FileNotFoundError as exc:
        raise RuntimeError(
            "LibreOffice is not installed or not found in PATH"
        ) from exc
    except subprocess.TimeoutExpired as exc:
        raise RuntimeError("LibreOffice conversion timed out") from exc

    if result.returncode != 0:
        logger.error("LibreOffice stderr: %s", result.stderr)
        raise RuntimeError(f"LibreOffice conversion failed: {result.stderr.strip()}")

    pdf_path = output_dir / (docx_path.stem + ".pdf")
    if not pdf_path.exists():
        raise RuntimeError(f"PDF not found after conversion: {pdf_path}")

    logger.info("PDF generated: %s", pdf_path)
    return pdf_path
